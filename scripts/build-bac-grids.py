"""Build the shared AFL1 catalogue from local school DOCX sources (python-docx).

Mappings/bands visually checked against the PDFs. Do not fetch national substitutes.
The three explicit local clarifications are recorded with the resulting rubric.
"""
from pathlib import Path
from docx import Document
import json, re
ROOT=Path(__file__).resolve().parents[1]
def row(id,t,r):
    cells=Document(ROOT/'bac_referentiels'/f'{id}.docx').tables[t].rows[r].cells
    result=[]; previous=None
    for c in cells:
        if c._tc is previous: continue
        previous=c._tc
        result.append(re.sub(r'\n{3,}','\n\n',c.text).strip())
    return result
def texts(id,t,r): return row(id,t,r)[1:5]
def criterion(title,maxi,descriptions,bands,**kw):
    assert len(descriptions)==4
    return dict(title=title,max=maxi,degrees=[dict(text=text,range=band) for text,band in zip(descriptions,bands)],**kw)
def spec(id,label,criteria,rule='',**kw):
    return dict(id=id,label=label,version=1,scale=12,source=id+'.docx',criteria=criteria,rule=rule,**kw)
band8=[[0,1.75],[2,3.75],[4,5.75],[6,8]]
band4=[[0,.75],[1,1.75],[2,2.75],[3,4]]
specs=[]
specs.append(spec('volley_ball','Volley-ball',[
 criterion('Utiliser et maîtriser les techniques de frappes adaptées à la situation et à l’intention de jeu',7,texts('volley_ball',4,2),[[.5,1.5],[2,3],[3.5,5],[5.5,7]],adjust='volley'),
 criterion('S’intégrer dans une organisation collective : placement et circulation du ballon',5,[s.replace('2,5--------------------------------3,','2,5 — 3,5 (borne confirmée par l’équipe)') for s in texts('volley_ball',4,4)],[[0,.5],[1,2],[2.5,3.5],[4,5]],adjust='volley')],
 'Proposition locale 2023. La note de chaque critère est multipliée par le coefficient de gain des matchs du degré choisi.',
 fields=[dict(key='outcome',label='Bilan des matchs',options=[['loss','Majorité de défaites'],['win','Majorité de victoires ou égalité']])],
 clarifications=['Borne haute du degré 3 collectif : 3,5 points, confirmée par le professeur le 15/10/2026.']))
bad1=texts('badminton',3,4);bad2=texts('badminton',3,6)
badchoices=lambda ns:[dict(label=label,value=n) for label,n in zip(['Plus de défaites','Autant de défaites que de victoires','Plus de victoires'],ns)]
specs.append(spec('badminton','Badminton',[
 criterion('Actions techniques d’attaque et de défense en relation avec le projet de jeu',8,bad1,band8,choices=list(map(badchoices,[[.5,1,1.75],[2,3,3.75],[4,5,5.75],[6,7,8]]))),
 criterion('Faire des choix au regard de l’analyse du rapport de force',4,bad2,band4,choices=list(map(badchoices,[[.25,.5,.75],[1,1.5,1.75],[2,2.5,2.75],[3,3.5,4]])))],
 'Choisir le degré, puis le bilan des matchs : la note est celle de la case du référentiel.',
 clarifications=['Tactique, degré 4, majorité de victoires : 4 points, confirmés par le professeur le 15/10/2026.']))
tt1=texts('tennis_de_table',2,3)
tt2=[a+'\n'+b+'\n'+c for a,b,c in zip(texts('tennis_de_table',2,13),texts('tennis_de_table',2,14),['Peu de projection et de choix tactiques à venir.','Il parvient à retenir quelques projections pour la période d’opposition à venir.','Les choix stratégiques sont efficaces et se traduisent en points gagnants ou fautes adverses.','Mise en place de projets tactiques à partir de son service pour placer un coup fort.'])]
specs.append(spec('tennis_de_table','Tennis de table',[
 criterion('Techniques de jeu : vitesse, hauteur et direction des trajectoires',8,tt1,[[0,2],[2.5,4],[4.5,6],[6.5,8]]),
 criterion('Pertinence des choix stratégiques',4,tt2,[[0,1],[1,2],[2,3],[3,4]])],
 'Après positionnement dans le degré, ajuster le curseur selon la proportion des oppositions gagnées individuellement : 0 % en bas, 100 % en haut de la fourchette. Le référentiel ne fixe pas d’interpolation automatique.'))
acro=[]
for col in range(1,5):
    fragments=[]
    for r in range(2,8):
        txt=Document(ROOT/'bac_referentiels/acrosport.docx').tables[2].rows[r].cells[col].text.strip()
        if txt and txt not in fragments: fragments.append(txt)
    acro.append('\n'.join(fragments))
specs.append(spec('acrosport','Acrosport',[
 criterion('Réaliser et maîtriser des formes corporelles complexes (individuel)',8,acro,[[0,2],[2.25,4],[4.25,6],[6.25,8]],adjust='acro'),
 criterion('Composer et présenter un enchaînement esthétique / acrobatique (collectif)',4,texts('acrosport',2,13),[[0,1],[1.25,2],[2.25,3],[3.25,4]])],
 'Quatre figures observées dans les rôles de porteur et voltigeur. Coefficient = somme des difficultés : A 0,10 ; B 0,15 ; C 0,20 ; D 0,25. Note individuelle = note du degré × coefficient − 0,5 par exigence manquante (minimum 0, maximum 8). Le critère collectif reste individuel dans la saisie : aucune copie implicite aux autres élèves.',
 fields=[*[dict(key=f'figure{i}',label=f'Figure {i+1}',options=[[s,s+' · '+str(n)] for s,n in [('A',.1),('B',.15),('C',.2),('D',.25)]]) for i in range(4)],dict(key='missing',label='Exigences manquantes',options=[[str(i),str(i)] for i in range(17)])],
 clarifications=['Ordre confirmé par le professeur le 15/10/2026 : multiplication, puis pénalités.']))
specs.append(spec('gymnastique','Gymnastique',[
 criterion('Réaliser avec maîtrise les éléments gymniques difficiles et adéquats',8,texts('gymnastique',2,5),[[.5,1.5],[2,3.5],[4,5],[5.5,7]],adjust='gym'),
 criterion('Composer et présenter un enchaînement esthétique / acrobatique',4,texts('gymnastique',2,7),[[.5,1],[1.5,2],[2.5,3],[3.5,4]])],
 'Six éléments, quatre familles et deux niveaux de difficulté. Le coefficient s’applique à la réalisation /8. Le référentiel autorise un arrondi : il reste un choix explicite du jury, désactivé par défaut.',
 fields=[dict(key='difficulty',label='Difficulté de l’enchaînement',options=[['0.9','Somme 2,6–3,4 · coefficient 0,9'],['1','Somme 3,6–4,6 · coefficient 1'],['1.1','Somme 4,8–5,8 · coefficient 1,1']]),dict(key='round',label='Arrondi du critère de réalisation',options=[['no','Sans arrondi'],['yes','Arrondir à l’entier (choix du jury)']],default='no')]))
points=[0,.25,.5,.75,1,1.25,1.5,1.75,2,2.25,2.5,2.75,3,3.2,3.5,3.75,4,4.2,4.5,4.75,5,5.2,5.5,5.75,6]
girls=[6.6,7.1,7.6,7.8,8.1,8.3,8.6,8.8,9.1,9.4,9.6,9.8,10.1,10.3,10.6,10.9,11.2,11.5,11.8,12.1,12.4,12.7,12.9,13.2,13.5]
boys=[7.8,8.6,9.4,9.7,10,10.3,10.7,11,11.4,11.7,12,12.3,12.5,12.8,13.1,13.4,13.7,14,14.4,14.8,15.2,15.6,16,16.5,17]
parts=[range(0,6),range(6,12),range(12,18),range(18,25)]
perf=[[dict(label=f'Filles {girls[i]} m · Garçons {boys[i]} m',value=points[i]) for i in inds] for inds in parts]
gaps=[[dict(label=f'Écart : {125-5*i} cm',value=points[i]) for i in inds] for inds in parts]
pt=texts('pentabond',2,7)[1:]
# This row has six logical cells (title, gap heading, four descriptions).
pt=row('pentabond',2,7)[2:6]
specs.append(spec('pentabond','Pentabond',[
 criterion('Meilleure performance réalisée',6,['Performance maximale, jusqu’à six essais. Premier bond obligatoirement en cloche-pied.\nSélectionner la ligne du barème filles / garçons correspondant à la meilleure performance validée.']*4,[[0,1.25],[1.5,2.75],[3,4.2],[4.5,6]],choices=perf),
 criterion('Efficacité technique · écart entre les deux meilleures performances',6,pt,[[0,1.25],[1.5,2.75],[3,4.2],[4.5,6]],choices=gaps)],
 'Sélection directe des lignes du barème : aucune interpolation non prévue n’est appliquée. Conserver les mesures réelles dans les observations. Un essai nul n’est pas une performance à zéro. Les valeurs 3,2 / 4,2 / 5,2 sont celles du document.'))
swim=[]
matrix=[[[.75,1,1.5],[.5,.75,1.25],[.5,.75,1.25],[.25,.5,1]],[[2,2.5,3],[1.75,2,2.5],[1.75,2,2.5],[1.5,1.75,2]],[[3.5,4,4.5],[3,3.5,4],[3,3.5,4],[2.5,3,3.5]],[[5,5.5,6],[4.5,5,5.5],[4.5,5,5.5],[4,4.5,5]]]
for d in range(4):
    swim.append([dict(label=f'{p} · {t} · obstacles {diff}',value=n) for (p,t),ns in zip([('P1 200 m','Temps 1'),('P1 200 m','Temps 2'),('P2 150 m','Temps 1'),('P2 150 m','Temps 2')],matrix[d]) for diff,n in zip(['A (3)','B (4)','C (6+)'],ns)])
tows=[[[.25,.5,1],[.5,1,1.5]],[[1.75,2,2.5],[2,2.5,3]],[[3.25,3.5,4],[3.5,4,4.5]],[[4.75,5,5.5],[5,5.5,6]]]
tow=[[dict(label=f'{man} · {dist}',value=n) for man,ns in zip(['Enfant','Adulte'],ds) for dist,n in zip(['A : 15–19 m','B : 20–24 m','C : 25–30 m'],ns)] for ds in tows]
times=['Filles P1 : T2 7:00 / T1 6:30 ; P2 : 5:50 / 5:20. Garçons P1 : 5:45 / 5:15 ; P2 : 4:50 / 4:20.', 'Filles P1 : T2 6:10 / T1 5:50 ; P2 : 5:05 / 4:40. Garçons P1 : 5:04 / 4:55 ; P2 : 4:19 / 4:00.', 'Filles P1 : T2 5:39 / T1 5:10 ; P2 : 4:29 / 4:00. Garçons P1 : 4:30 / 4:15 ; P2 : 3:40 / 3:20.', 'Filles P1 : T2 5:00 / T1 4:30 ; P2 : 3:50 / 3:20. Garçons P1 : 4:04 / 3:45 ; P2 : 3:10 / 2:50.']
specs.append(spec('sauvetage','Sauvetage',[
 criterion('Adapter le déplacement et franchir efficacement les obstacles',6,[a+'\n\nRepères chronométriques (minutes:secondes) :\n'+b for a,b in zip(texts('sauvetage',1,5),times)],[[.25,1.5],[1.5,3],[2.5,4.5],[4,6]],choices=swim),
 criterion('Maîtriser le déplacement avec le mannequin remorqué',6,texts('sauvetage',1,6),[[.25,1.5],[1.75,3],[3.25,4.5],[4.75,6]],choices=tow)],
 'Nage : temps réalisé + 5 secondes par pénalité. Après lecture des repères du degré, sélectionner le parcours, le repère Temps 1/2 et la difficulté. Le choix du jury est explicite : pas d’interpolation des temps. Remorquage : une minute maximum.'))
specs.append(spec('course_en_duree','Course en durée',[
 criterion('Produire la séance',8,texts('course_en_duree',3,5),band8),
 criterion('Analyser la séance',4,texts('course_en_duree',3,6),band4)],
 'Séance de 30 à 45 minutes. Évaluer l’échauffement, l’engagement, l’adaptation de la foulée et l’analyse des ressentis selon le projet choisi.'))
specs.append(spec('musculation','Musculation',[
 criterion('Produire : gestuelle, postures, charge et choix des exercices',8,texts('musculation',1,3),[[0,2.5],[3,4.5],[5,6.5],[7,8]]),
 criterion('Analyser : ressentis et régulation de la charge',4,texts('musculation',1,4),[[0,1],[1.5,2],[2.5,3],[3.5,4]])],
 'Séance de 60 minutes. Thème endurance, volume ou puissance. Deux groupes musculaires minimum, deux à trois exercices par groupe.'))
catalog={s['id']:s for s in specs}
assert len(catalog)==9
data=json.dumps(catalog,ensure_ascii=False,indent=2)+'\n'
(ROOT/'bac_referentiels/afl1-grids.json').write_text(data,encoding='utf-8')
(ROOT/'bac-rubrics.js').write_text('/* Generated by scripts/build-bac-grids.py; local school references. */\n(function(r){const specs='+data.strip()+';r.EpsBacRubrics=specs;if(typeof module==="object"&&module.exports)module.exports=specs;})(globalThis);\n',encoding='utf-8')
android=ROOT.parent/'installation-manager-current/app/src/main/assets/bac_referentiels/afl1-grids.json'
android.write_text(data,encoding='utf-8')
